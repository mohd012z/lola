"""High-fidelity helpers for Lola Office -> interactive HTML."""
from __future__ import annotations
import base64, html, io
from pathlib import Path
from lola_office_html_styles import xlsx_style, xlsx_images, docx_headers_footers, hyperlink_map
from lola_office_layout import word_body_in_order, excel_layout, excel_colgroup

def data_uri(blob,mime):
    return "data:"+mime+";base64,"+base64.b64encode(blob).decode("ascii")

def docx_rich(src):
    from docx import Document
    doc=Document(src)
    return docx_headers_footers(doc)+word_body_in_order(doc,data_uri)

"""High-fidelity helpers for Lola Office -> interactive HTML."""
from __future__ import annotations
import base64, html, io
from pathlib import Path
from lola_office_html_styles import xlsx_style, xlsx_images, docx_headers_footers, hyperlink_map
from lola_office_layout import word_body_in_order, excel_layout, excel_colgroup

def data_uri(blob,mime):
    return "data:"+mime+";base64,"+base64.b64encode(blob).decode("ascii")

def docx_rich(src):
    from docx import Document
    doc=Document(src);parts=[docx_headers_footers(doc)]
    for p in doc.paragraphs:
        style=(p.style.name if p.style else "")
        level=None
        if style.lower().startswith("heading"):
            try:level=max(1,min(6,int(style.split()[-1])))
            except:level=2
        runs=[]
        for r in p.runs:
            t=html.escape(r.text)
            if r.bold:t="<strong>"+t+"</strong>"
            if r.italic:t="<em>"+t+"</em>"
            if r.underline:t="<u>"+t+"</u>"
            runs.append(t)
        text="".join(runs)
        for label,url in hyperlink_map(p):
            safe_url=html.escape(url,quote=True)
            text=text.replace(html.escape(label),"<a href='"+safe_url+"' rel='noopener noreferrer'>"+html.escape(label)+"</a>",1)
        if text:parts.append(("<h%d>%s</h%d>"%(level,text,level)) if level else "<p>"+text+"</p>")
    # Embedded images are exported as data URIs.
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            part=rel.target_part;mime=part.content_type
            parts.append("<figure><img loading='lazy' src='"+data_uri(part.blob,mime)+"'></figure>")
    return "".join(parts)

def xlsx_rich(src):
    from openpyxl import load_workbook
    wb=load_workbook(src,data_only=False);parts=[]
    for ws in wb.worksheets:
        merged={str(r) for r in ws.merged_cells.ranges};skip=set();rows=[]
        for r in range(1,ws.max_row+1):
            cells=[]
            for c in range(1,ws.max_column+1):
                if (r,c) in skip:continue
                cell=ws.cell(r,c);rs=cs=1
                for rng in ws.merged_cells.ranges:
                    if rng.min_row==r and rng.min_col==c:
                        rs=rng.max_row-rng.min_row+1;cs=rng.max_col-rng.min_col+1
                        for rr in range(r,r+rs):
                            for cc in range(c,c+cs):
                                if (rr,cc)!=(r,c):skip.add((rr,cc))
                        break
                attrs=(" rowspan='%d'"%rs if rs>1 else "")+(" colspan='%d'"%cs if cs>1 else "")
                v="" if cell.value is None else str(cell.value)
                style=xlsx_style(cell)
                cells.append("<td"+attrs+" data-cell='"+cell.coordinate+"' style='"+html.escape(style,quote=True)+"'>"+html.escape(v)+"</td>")
            rd=ws.row_dimensions[r]\n            rstyle=("height:%spt"%rd.height if rd.height else "")+(";display:none" if rd.hidden else "")\n            rows.append("<tr style='"+html.escape(rstyle,quote=True)+"'>"+"".join(cells)+"</tr>")
        parts.append("<section class='sheet' data-sheet='"+html.escape(ws.title,quote=True)+"'><h2>"+html.escape(ws.title)+"</h2><div class='table-wrap' data-freeze='"+html.escape(str(excel_layout(ws).get("freeze_panes") or ""),quote=True)+"'><table>"+excel_colgroup(ws)+"".join(rows)+"</table></div>"+xlsx_images(ws)+"</section>")
    return "".join(parts)

def interactive_shell(title,body):
    css="""body{margin:0;font-family:system-ui;background:#eef1f5;color:#161616}.toolbar{position:sticky;top:0;z-index:9;display:flex;gap:8px;padding:8px;background:#20242a;color:white}.toolbar input{flex:1;min-width:80px}.doc{max-width:1180px;margin:auto;background:white;padding:20px;min-height:100vh}.table-wrap{overflow:auto}table{border-collapse:collapse}td,th{border:1px solid #aaa;padding:6px;min-width:40px}img{max-width:100%;height:auto}.hit{outline:2px solid currentColor}@media print{.toolbar{display:none}.doc{padding:0}}"""
    js="""function q(){let s=document.getElementById('q').value.toLowerCase();document.querySelectorAll('.hit').forEach(x=>x.classList.remove('hit'));if(!s)return;for(const e of document.querySelectorAll('p,h1,h2,h3,h4,h5,h6,td'))if(e.textContent.toLowerCase().includes(s))e.classList.add('hit')}function z(n){let d=document.querySelector('.doc'),v=parseFloat(d.dataset.z||1);v=Math.max(.5,Math.min(2.5,v+n));d.dataset.z=v;d.style.zoom=v}"""
    return "<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>"+html.escape(title)+"</title><style>"+css+"</style></head><body><div class='toolbar'><button onclick='z(-.1)'>−</button><button onclick='z(.1)'>+</button><input id='q' oninput='q()' placeholder='Search document'><button onclick='print()'>Print / PDF</button></div><main class='doc'>"+body+"</main><script>"+js+"</script></body></html>"

def save_rich(src,dst,kind):
    body=docx_rich(src) if kind=="docx" else xlsx_rich(src)
    p=Path(dst);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(interactive_shell(Path(src).name,body),encoding="utf-8")
    return {"ok":True,"output":str(p),"mode":"interactive-rich-html"}
