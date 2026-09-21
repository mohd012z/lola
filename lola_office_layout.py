"""Document-order Word renderer and richer Excel layout metadata for Lola."""
from __future__ import annotations
import html

def _tag(el):return el.tag.rsplit("}",1)[-1]

def word_body_in_order(doc,data_uri_fn):
    out=[]
    rels=doc.part.rels
    for el in doc.element.body.iterchildren():
        tag=_tag(el)
        if tag=="p":
            texts=[];links={}
            for n in el.iter():
                nt=_tag(n)
                if nt=="t":texts.append(html.escape(n.text or ""))
                elif nt=="blip":
                    rid=n.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if rid in rels:
                        part=rels[rid].target_part
                        texts.append("<img loading='lazy' class='inline-image' src='"+data_uri_fn(part.blob,part.content_type)+"'>")
            txt="".join(texts)
            pstyle=""
            for n in el.iter():
                if _tag(n)=="pStyle":pstyle=n.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val","");break
            is_list=any(_tag(n)=="numPr" for n in el.iter())
            if txt:
                if pstyle.lower().startswith("heading"):
                    digits="".join(x for x in pstyle if x.isdigit());lv=max(1,min(6,int(digits or "2")))
                    out.append("<h%d>%s</h%d>"%(lv,txt,lv))
                elif is_list:out.append("<div class='list-item'>• "+txt+"</div>")
                else:out.append("<p>"+txt+"</p>")
        elif tag=="tbl":
            rows=[]
            for tr in (x for x in el if _tag(x)=="tr"):
                cells=[]
                for tc in (x for x in tr if _tag(x)=="tc"):
                    text="".join(html.escape(n.text or "") for n in tc.iter() if _tag(n)=="t")
                    cells.append("<td>"+text+"</td>")
                rows.append("<tr>"+"".join(cells)+"</tr>")
            out.append("<div class='table-wrap'><table>"+"".join(rows)+"</table></div>")
    return "".join(out)

def excel_layout(ws):
    cols={}
    for key,d in ws.column_dimensions.items():
        if d.width:cols[key]={"width":d.width,"hidden":bool(d.hidden)}
    rows={}
    for key,d in ws.row_dimensions.items():
        if d.height or d.hidden:rows[str(key)]={"height":d.height,"hidden":bool(d.hidden)}
    pane=str(ws.freeze_panes) if ws.freeze_panes else None
    return {"columns":cols,"rows":rows,"freeze_panes":pane,"sheet_hidden":ws.sheet_state!="visible"}

def excel_colgroup(ws):
    from openpyxl.utils import get_column_letter
    out=["<colgroup>"]
    for c in range(1,ws.max_column+1):
        key=get_column_letter(c);d=ws.column_dimensions[key]
        style=[]
        if d.width:style.append("width:%.2fch"%d.width)
        if d.hidden:style.append("display:none")
        out.append("<col style='"+html.escape(";".join(style),quote=True)+"'>")
    out.append("</colgroup>")
    return "".join(out)
