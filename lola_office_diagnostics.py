"""Detailed missing-content diagnostics for Office -> HTML conversion."""
from __future__ import annotations
import html, json, re
from pathlib import Path

def _sheet_source(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False)
    out=[]
    for ws in wb.worksheets:
        out.append({"name":ws.title,"cells":sum(1 for row in ws.iter_rows() for c in row if c.value is not None),
                    "images":len(getattr(ws,"_images",[])),"charts":len(getattr(ws,"_charts",[])),
                    "merged":len(ws.merged_cells.ranges),"state":ws.sheet_state})
    return out

def _sheet_html(text):
    out={}
    pat=re.compile(r"<section[^>]*class=['\"]sheet['\"][^>]*data-sheet=['\"]([^'\"]+)['\"][^>]*>(.*?)</section>",re.I|re.S)
    for name,body in pat.findall(text):
        out[html.unescape(name)]={"cells":len(re.findall(r"<td(?:\s|>)",body,re.I)),
            "images":len(re.findall(r"<img(?:\s|>)",body,re.I)),
            "charts":len(re.findall(r"class=['\"]chart-card",body,re.I))}
    return out

def diagnose(source,output):
    src=Path(source);text=Path(output).read_text(encoding="utf-8",errors="replace")
    issues=[]
    if src.suffix.lower()==".xlsx":
        rendered=_sheet_html(text)
        for sh in _sheet_source(src):
            got=rendered.get(sh["name"])
            if got is None:
                issues.append({"severity":"error","location":"sheet:"+sh["name"],"issue":"sheet missing from HTML"})
                continue
            for key in ("cells","images","charts"):
                if got[key] < sh[key]:
                    issues.append({"severity":"warning","location":"sheet:"+sh["name"],"issue":key+" coverage low",
                                   "expected":sh[key],"actual":got[key]})
    elif src.suffix.lower()==".docx":
        from docx import Document
        d=Document(src)
        st=len(d.tables);ht=len(re.findall(r"<table(?:\s|>)",text,re.I))
        if ht<st:issues.append({"severity":"error","location":"document","issue":"table coverage low","expected":st,"actual":ht})
        si=sum(1 for r in d.part.rels.values() if "image" in r.reltype);hi=len(re.findall(r"<img(?:\s|>)",text,re.I))
        if hi<si:issues.append({"severity":"warning","location":"document","issue":"image coverage low","expected":si,"actual":hi})
    elif src.suffix.lower()==".pdf":
        import fitz
        d=fitz.open(src);pages=set(int(x) for x in re.findall(r"data-page=['\"](\d+)",text))
        for n in range(1,len(d)+1):
            if n not in pages:issues.append({"severity":"error","location":"page:"+str(n),"issue":"page missing from HTML"})
    return {"issues":issues,"issue_count":len(issues),"errors":sum(i["severity"]=="error" for i in issues),
            "warnings":sum(i["severity"]=="warning" for i in issues),
            "status":"attention" if issues else "structurally-clear",
            "note":"Diagnostics identify structural omissions; they are not a pixel-level visual comparison."}

def save_report(report,path):
    p=Path(path);p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding="utf-8")
    return str(p)
