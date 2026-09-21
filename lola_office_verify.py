"""Post-conversion coverage verification for Lola Office -> HTML."""
from __future__ import annotations
import re
from pathlib import Path

def source_inventory(path):
    p=Path(path);ext=p.suffix.lower()
    if ext==".docx":
        from docx import Document
        d=Document(p)
        images=sum(1 for r in d.part.rels.values() if "image" in r.reltype)
        return {"kind":"docx","paragraphs":sum(1 for x in d.paragraphs if x.text.strip()),
                "tables":len(d.tables),"images":images,"sections":len(d.sections)}
    if ext==".xlsx":
        from openpyxl import load_workbook
        wb=load_workbook(p,data_only=False)
        return {"kind":"xlsx","sheets":len(wb.worksheets),
                "nonempty_cells":sum(sum(1 for row in ws.iter_rows() for c in row if c.value is not None) for ws in wb.worksheets),
                "images":sum(len(getattr(ws,"_images",[])) for ws in wb.worksheets),
                "charts":sum(len(getattr(ws,"_charts",[])) for ws in wb.worksheets),
                "merged_ranges":sum(len(ws.merged_cells.ranges) for ws in wb.worksheets)}
    if ext==".pdf":
        import fitz
        d=fitz.open(p)
        return {"kind":"pdf","pages":len(d),"text_pages":sum(1 for pg in d if pg.get_text("text").strip())}
    return {"kind":ext.lstrip(".") or "unknown","bytes":p.stat().st_size}

def html_inventory(path):
    s=Path(path).read_text(encoding="utf-8",errors="replace")
    return {"paragraphs":len(re.findall(r"<p(?:\s|>)",s,re.I)),
            "tables":len(re.findall(r"<table(?:\s|>)",s,re.I)),
            "images":len(re.findall(r"<img(?:\s|>)",s,re.I)),
            "sheets":len(re.findall(r"<section[^>]+class=['\"]sheet",s,re.I)),
            "pdf_pages":len(re.findall(r"<section[^>]+class=['\"]pdf-page",s,re.I)),
            "cells":len(re.findall(r"<td(?:\s|>)",s,re.I)),
            "charts":len(re.findall(r"class=['\"]chart-card",s,re.I))}

def verify(source,output):
    src=source_inventory(source);dst=html_inventory(output);checks=[]
    mapping={"tables":"tables","images":"images","sheets":"sheets","charts":"charts","pages":"pdf_pages"}
    for sk,dk in mapping.items():
        if sk in src:
            expected=src[sk];actual=dst.get(dk,0)
            checks.append({"feature":sk,"expected":expected,"actual":actual,"covered":actual>=expected})
    if src.get("kind")=="xlsx" and "nonempty_cells" in src:
        checks.append({"feature":"nonempty_cells","expected":src["nonempty_cells"],"actual":dst["cells"],
                       "covered":dst["cells"]>=src["nonempty_cells"]})
    covered=sum(1 for x in checks if x["covered"])
    return {"source":src,"html":dst,"checks":checks,
            "coverage_checks_passed":covered,"coverage_checks_total":len(checks),
            "complete":bool(checks) and covered==len(checks),
            "note":"Structural coverage only; passing does not prove pixel-identical rendering."}
