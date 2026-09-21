"""Built-in office conversion execution engine for Lola.

Local, user-selected documents only. Uses Python libraries for editable structures
and LibreOffice when a renderer is required. PDF reconstruction is best-effort.
"""
from __future__ import annotations
import csv, shutil, subprocess, tempfile
from pathlib import Path

def capabilities():
    mods={}
    for name in ("docx","openpyxl","fitz","pdfplumber"):
        try:__import__(name);mods[name]=True
        except ImportError:mods[name]=False
    mods["libreoffice"]=bool(shutil.which("libreoffice") or shutil.which("soffice"))
    return mods

def _office_export(src,out_dir,fmt):
    exe=shutil.which("libreoffice") or shutil.which("soffice")
    if not exe:return {"ok":False,"error":"LibreOffice/soffice not available"}
    d=Path(out_dir);d.mkdir(parents=True,exist_ok=True)
    p=subprocess.run([exe,"--headless","--convert-to",fmt,"--outdir",str(d),str(src)],
                     capture_output=True,text=True,timeout=180,check=False)
    return {"ok":p.returncode==0,"returncode":p.returncode,"stdout":p.stdout[-8000:],"stderr":p.stderr[-8000:]}

def docx_to_xlsx(src,dst):
    from docx import Document
    from openpyxl import Workbook
    doc=Document(src);wb=Workbook();wb.remove(wb.active)
    if not doc.tables:
        ws=wb.create_sheet("Document")
        for i,p in enumerate(doc.paragraphs,1):ws.cell(i,1,p.text)
    for ti,t in enumerate(doc.tables,1):
        ws=wb.create_sheet(("Table "+str(ti))[:31])
        for r,row in enumerate(t.rows,1):
            for c,cell in enumerate(row.cells,1):ws.cell(r,c,cell.text)
    Path(dst).parent.mkdir(parents=True,exist_ok=True);wb.save(dst)
    return {"ok":True,"output":str(dst),"tables":len(doc.tables)}

def xlsx_to_docx(src,dst):
    from openpyxl import load_workbook
    from docx import Document
    wb=load_workbook(src,data_only=False,read_only=True);doc=Document()
    for si,ws in enumerate(wb.worksheets):
        if si:doc.add_page_break()
        doc.add_heading(ws.title,level=1)
        rows=list(ws.iter_rows(values_only=True))
        if not rows:continue
        cols=max(len(r) for r in rows);table=doc.add_table(rows=len(rows),cols=cols)
        for r,row in enumerate(rows):
            for c,v in enumerate(row):table.cell(r,c).text="" if v is None else str(v)
    Path(dst).parent.mkdir(parents=True,exist_ok=True);doc.save(dst)
    return {"ok":True,"output":str(dst),"sheets":len(wb.worksheets)}

def csv_to_xlsx(src,dst):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.title="Data"
    with open(src,newline="",encoding="utf-8-sig",errors="replace") as f:
        for row in csv.reader(f):ws.append(row)
    Path(dst).parent.mkdir(parents=True,exist_ok=True);wb.save(dst)
    return {"ok":True,"output":str(dst)}

def xlsx_to_csv(src,dst,sheet=None):
    from openpyxl import load_workbook
    wb=load_workbook(src,data_only=True,read_only=True);ws=wb[sheet] if sheet else wb.active
    Path(dst).parent.mkdir(parents=True,exist_ok=True)
    with open(dst,"w",newline="",encoding="utf-8") as f:
        w=csv.writer(f)
        for row in ws.iter_rows(values_only=True):w.writerow(["" if v is None else v for v in row])
    return {"ok":True,"output":str(dst),"sheet":ws.title}

def pdf_to_docx(src,dst):
    import fitz
    from docx import Document
    doc=Document();pdf=fitz.open(src)
    for pi,page in enumerate(pdf):
        if pi:doc.add_page_break()
        blocks=page.get_text("blocks")
        for b in sorted(blocks,key=lambda x:(round(x[1],1),x[0])):
            text=(b[4] or "").strip()
            if text:doc.add_paragraph(text)
    Path(dst).parent.mkdir(parents=True,exist_ok=True);doc.save(dst)
    return {"ok":True,"output":str(dst),"pages":len(pdf),"mode":"editable-text-reconstruction"}

def pdf_to_xlsx(src,dst):
    import pdfplumber
    from openpyxl import Workbook
    wb=Workbook();wb.remove(wb.active);tables=0
    with pdfplumber.open(src) as pdf:
        for pi,page in enumerate(pdf,1):
            for ti,table in enumerate(page.extract_tables() or [],1):
                ws=wb.create_sheet(("P%d T%d"%(pi,ti))[:31]);tables+=1
                for row in table:ws.append(["" if v is None else v for v in row])
    if not wb.worksheets:wb.create_sheet("No tables detected")
    Path(dst).parent.mkdir(parents=True,exist_ok=True);wb.save(dst)
    return {"ok":True,"output":str(dst),"tables":tables}

def convert(src,dst):
    src=Path(src);dst=Path(dst);s=src.suffix.lower();d=dst.suffix.lower()
    if s==".docx" and d==".xlsx":return docx_to_xlsx(src,dst)
    if s==".xlsx" and d==".docx":return xlsx_to_docx(src,dst)
    if s==".csv" and d==".xlsx":return csv_to_xlsx(src,dst)
    if s==".xlsx" and d==".csv":return xlsx_to_csv(src,dst)
    if s==".pdf" and d==".docx":return pdf_to_docx(src,dst)
    if s==".pdf" and d==".xlsx":return pdf_to_xlsx(src,dst)
    if d==".pdf":return _office_export(src,dst.parent,"pdf")
    # Office-to-office legacy/ODF conversions are delegated to LibreOffice.
    fmt=d.lstrip(".")
    if s in (".doc",".docx",".odt",".rtf",".xls",".xlsx",".ods",".csv") and fmt:
        return _office_export(src,dst.parent,fmt)
    return {"ok":False,"error":"No safe built-in conversion route","source":s,"target":d}
